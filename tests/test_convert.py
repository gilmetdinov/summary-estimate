"""Интеграционный тест конвертации (convert) на реальном template.xlsx."""
from __future__ import annotations

import os

import openpyxl
import pytest

from render import convert, default_report_name
from render.donor import DEFAULT_DONOR
from config import Config


def test_default_report_name(raw_xlsx):
    from smeta_parser import parse_workbook
    smetas = parse_workbook(raw_xlsx)
    name = default_report_name(smetas)
    assert name.startswith("Сводник")
    assert name.endswith(".xlsx")


@pytest.mark.skipif(not os.path.exists(DEFAULT_DONOR), reason="нет template.xlsx")
def test_convert_produces_valid_workbook(raw_xlsx, tmp_path):
    out = tmp_path / "result.xlsx"
    convert(raw_xlsx, str(out), DEFAULT_DONOR, Config())
    assert out.exists()

    wb = openpyxl.load_workbook(str(out))
    # сводная идёт первым листом, смета — следующим
    assert "св.смета" in wb.sheetnames
    assert len(wb.sheetnames) >= 2


@pytest.mark.skipif(not os.path.exists(DEFAULT_DONOR), reason="нет template.xlsx")
def test_convert_raises_on_empty_input(tmp_path):
    empty = tmp_path / "empty.xlsx"
    openpyxl.Workbook().save(str(empty))
    with pytest.raises(ValueError):
        convert(str(empty), str(tmp_path / "out.xlsx"), DEFAULT_DONOR, Config())
