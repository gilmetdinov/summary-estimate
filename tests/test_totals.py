"""Консистентность итогов: формулы раздела/НДС/ВСЕГО дают значения выгрузки.

Синтетика зеркалит семантику ГС: E подстрок разбивки в сумме дают E позиции,
сумма позиций — субтотал выгрузки. Проверяем, что пересобранный лист
пересчитывается в те же числа (эмуляция SUM/ROUND без Excel).
"""
from __future__ import annotations

import re

import openpyxl
import pytest

from config import Config
from render import convert
from render.donor import DEFAULT_DONOR


def _sheet(rows: list[tuple], title: str):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = title
    for r, row in enumerate(rows, start=1):
        for c, val in enumerate(row, start=1):
            if val is not None:
                ws.cell(row=r, column=c).value = val
    return wb


def _header(number: str, object_name: str) -> list[tuple]:
    return [
        (f"СМЕТА № {number}", None, None, None, None),
        (object_name, None, None, None, None),
        ("на проектные работы", None, None, None, None),
        ("Составлена в уровне цен на 1 квартал 2025 г.", None, None, None, None),
        ("Заказчик", None, None, None, None),
        (None, "ООО «Заказчик»", None, None, None),
        ("Проектная организация", None, None, None, None),
        (None, "ООО «Проектная организация»", None, None, None),
        ("№ пп", None, None, None, None),
        ("1", "2", "3", "4", "5"),
    ]


def _totals(sub: float, vat: float, total: float) -> list[tuple]:
    return [
        (None, "Итоги по смете:", None, None, None),
        (None, "     Итого Поз. 1-2", None, None, sub),
        (None, "     НДС 22%", None, None, vat),
        (None, "     ВСЕГО по смете", None, None, total),
    ]


def _eval_chain(ws) -> dict[str, float]:
    """Собирает числа колонки E и вычисляет формулы SUM/ROUND по цепочке."""
    vals: dict[int, float] = {}
    for r in range(1, ws.max_row + 1):
        e = ws.cell(row=r, column=5).value
        if isinstance(e, (int, float)):
            vals[r] = float(e)

    def ev(f: str) -> float:
        for m in sorted(set(re.findall(r"E(\d+)", f)), key=int, reverse=True):
            f = f.replace(f"E{m}", repr(vals.get(int(m), 0.0)))
        f = f.lstrip("=").replace("ROUND", "round").replace("SUM", "_s")
        return eval(f, {"_s": lambda *a: sum(a), "round": round})  # noqa: S307

    out: dict[str, float] = {}
    for r in range(1, ws.max_row + 1):
        b = str(ws.cell(row=r, column=2).value or "")
        f = ws.cell(row=r, column=5).value
        if not (isinstance(f, str) and f.startswith("=")):
            continue
        if "Итого по разделу" in b:
            vals[r] = ev(f)
            out.setdefault("sec", vals[r])
        elif "Итого Поз" in b:
            vals[r] = ev(f)
            out["sub"] = vals[r]
        elif "НДС" in b:
            vals[r] = ev(f)
            out["nds"] = vals[r]
        elif "ВСЕГО" in b:
            vals[r] = ev(f)
            out["vs"] = vals[r]
    return out


def _labels(ws) -> dict[str, str]:
    """label-строка -> формула E (для точных проверок подписей)."""
    out = {}
    for r in range(1, ws.max_row + 1):
        b = str(ws.cell(row=r, column=2).value or "").strip()
        if b.startswith(("Итого по разделу", "Итого Поз", "НДС", "ВСЕГО")):
            out[b] = ws.cell(row=r, column=5).value
    return out


@pytest.mark.skipif(
    __import__("os").path.exists(DEFAULT_DONOR) is False, reason="нет template.xlsx")
def test_totals_chain_matches_source(tmp_path):
    rows = _header("12.01.01", "Строительство объекта — пример очень длинного наименования стройки") + [
        ("Раздел 1. Проектная документация", None, None, None, None),
        ("1", "Позиция один", "МНЗ п.1", "расчёт 1", 100.0),
        (None, None, "ПЗ 0,5%;", None, 40.0),
        (None, None, "ТКР 75,5%;", None, 60.0),
        ("2", "Позиция два", "МНЗ п.2", "расчёт 2", 50.0),
    ] + _totals(150.0, 33.0, 183.0)
    src = tmp_path / "raw.xlsx"
    _sheet(rows, "смета 12.01.01").save(str(src))
    out = tmp_path / "result.xlsx"
    convert(str(src), str(out), DEFAULT_DONOR, Config())

    wb = openpyxl.load_workbook(str(out), data_only=False)
    ws = wb[wb.sheetnames[1]]
    got = _eval_chain(ws)
    assert got == pytest.approx({"sec": 150.0, "sub": 150.0, "nds": 33.0, "vs": 183.0})

    labs = _labels(ws)
    assert "Итого по разделу 1 Проектная документация" in labs
    assert str(labs["НДС 22%"]).startswith("=ROUND(")
    assert str(labs["ВСЕГО по смете"]).startswith("=")

    # сумма прописью — от ВСЕГО с НДС
    propis = str(ws.cell(row=1, column=1).value or "")
    for r in range(1, ws.max_row + 1):
        v = ws.cell(row=r, column=1).value
        if isinstance(v, str) and v.startswith('=CONCATENATE("Всего по смете'):
            propis = v
    assert "Всего по смете с НДС" in propis

    # печать: сквозные строки и вписывание по ширине
    assert ws.print_title_rows
    assert ws.page_setup.fitToWidth == 1
    svod = wb["св.смета"]
    assert str(svod.print_area).endswith(f"$G${svod.max_row}") or ":$G$" in str(svod.print_area)


@pytest.mark.skipif(
    __import__("os").path.exists(DEFAULT_DONOR) is False, reason="нет template.xlsx")
def test_multi_section_labels(tmp_path):
    rows = _header("12.03", "Изыскания по объекту — пример очень длинного наименования стройки") + [
        ("Раздел 1. Полевые работы", None, None, None, None),
        ("1", "Бурение", "МНЗ п.1", "расчёт 1", 200.0),
        ("Раздел 2. Камеральные работы", None, None, None, None),
        ("2", "Обработка", "МНЗ п.2", "расчёт 2", 60.0),
    ] + _totals(260.0, 57.2, 317.2)
    src = tmp_path / "raw.xlsx"
    _sheet(rows, "смета 12.03").save(str(src))
    out = tmp_path / "result.xlsx"
    convert(str(src), str(out), DEFAULT_DONOR, Config())

    wb = openpyxl.load_workbook(str(out), data_only=False)
    ws = wb[wb.sheetnames[1]]
    labs = _labels(ws)
    assert "Итого по разделу 1 Полевые работы" in labs
    assert "Итого по разделу 2 Камеральные работы" in labs
    got = _eval_chain(ws)
    assert got == pytest.approx({"sec": 200.0, "sub": 260.0, "nds": 57.2, "vs": 317.2})
