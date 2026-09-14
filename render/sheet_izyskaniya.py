"""
Рендер листа-заготовки изысканий (пустые строки по трём разделам, итог, подписи).
"""
from __future__ import annotations
from copy import copy

from openpyxl.utils import get_column_letter

from config import Config
from render.donor import Donor
from render.style import clean_quotes, merge_block, put


def render_izyskaniya_placeholder(wb, d: Donor, cfg: Config,
                                  used: set, project_code: str) -> str:
    """Генерирует лист-заготовку изысканий (3 раздела, пустые строки, итог)."""
    S = d.sm
    base = f"изыскания_{project_code}"[:31]
    title, i = base, 1
    while title in used:
        i += 1
        sfx = f"_{i}"
        title = f"{base[:31 - len(sfx)]}{sfx}"
    used.add(title)

    ws = wb.create_sheet(title=title)
    for col, w in {"A": 7.0, "B": 32.1, "C": 47.1, "D": 40.3, "E": 12.1}.items():
        ws.column_dimensions[col].width = w
    ws.column_dimensions["F"].width = 3
    ws.column_dimensions["F"].hidden = True
    ws.sheet_view.showGridLines = False

    ws.cell(row=1, column=6).value = project_code
    ws.cell(row=6, column=6).value = project_code

    put(ws, 1, 5, '=CONCATENATE("Приложение №3.",F1,".","3")', S["E1"], height=12.75)
    put(ws, 2, 5, "к договору", S["E2"], height=12.75)
    for r in (2, 3, 4):
        c = ws.cell(row=r, column=5)
        a = copy(c.alignment); a.wrap_text = False; c.alignment = a
    put(ws, 3, 5, "от ____.____._______г.", S["E2"], height=12.75)
    put(ws, 4, 5, "№_________________", S["E2"], height=12.75)
    put(ws, 6, 1, '=CONCATENATE("СМЕТА № ",F6,".","03")', S["A6"], c2=5, height=15.0)
    ws.row_dimensions[17].hidden = True
    put(ws, 7, 1, "на изыскательские работы", S["A7"], c2=5, height=15.0)
    put(ws, 8, 1, clean_quotes(cfg.izyskaniya.name) if cfg.izyskaniya else "",
        S["A8"], c2=5, height=25.5)
    put(ws, 9, 1, "(наименование стройки)", S["A9"], c2=5, height=15.0)
    put(ws, 11, 1, "Заказчик", S["A11"], height=15.0)
    put(ws, 12, 2, "", S["B12"], c2=5, height=15.0)
    put(ws, 13, 1, "(наименование организации)", S["A13"], c2=5, height=15.0)
    put(ws, 14, 1, "Проектная организация", S["A14"], height=15.0)
    put(ws, 15, 2, cfg.org_short, S["B15"], c2=5, height=15.0)
    put(ws, 16, 1, "(наименование организации)", S["A16"], c2=5, height=15.0)
    put(ws, 18, 1, "Составлена в уровне цен на ____ квартал 20__ г.", S["A18"], height=15.0)

    hdr = ["№ пп",
           "Наименование\nобъекта проектирования или вида проектных работ",
           "Наименование, номера глав, таблиц, параграфов и пунктов МНЗ на проектные работы",
           "Расчет стоимости", "Сметная стоимость, тыс. руб."]
    for c, txt in enumerate(hdr, start=1):
        put(ws, 20, c, txt, S[f"{get_column_letter(c)}20"], height=36.0)
    for c in range(1, 6):
        put(ws, 21, c, c, S["A21"], height=15.0)

    sections = [
        "Раздел 1. Полевые работы",
        "Раздел 2. Камеральные работы",
        "Раздел 3. Составление отчета",
    ]
    R = 22
    first_data = R + 1
    for sec_title in sections:
        put(ws, R, 1, sec_title, S["A22"], c2=5, height=15.0)
        R += 1
        # пустая строка-плейсхолдер (все 5 колонок с границами из A23..E23)
        for c in range(1, 6):
            put(ws, R, c, None, S[f"{get_column_letter(c)}23"])
        R += 2

    # R после цикла — позиция ЗА последней data-строкой (rows: секция, data, пропуск).
    # Для 3 разделов data-строки: 23, 26, 29; R=31 → last_data = R-2 = 29.
    last_data = R - 2
    put(ws, R, 1, None, S["A49"]); put(ws, R, 2, "Итоги по смете:", S["B50"], c2=3, height=15.0); put(ws, R, 4, None, S["D50"]); put(ws, R, 5, None, S["E51"], height=15.0)
    R += 1
    subtotal_addr = f"E{R}"
    put(ws, R, 1, None, S["A51"]); put(ws, R, 2, "     Итого Поз. 1-3", S["B51"], c2=3, height=15.0); put(ws, R, 4, None, S["D51"])
    sum_range = ",".join(f"E{r}" for r in range(first_data, last_data + 1, 3))
    put(ws, R, 5, f"=SUM({sum_range})", S["E51"], height=15.0)
    R += 2

    put(ws, R, 1, d.propis_for_smeta(subtotal_addr), S["A55"], c2=5, height=12.75)
    R += 1
    put(ws, R, 1, "(сумма прописью)", S["A56"], c2=5, height=12.75)
    R += 2

    put(ws, R, 1, "Проектная организация:", S["A57"], height=12.75); R += 1
    dir_block = f"{cfg.org_dir_title1} {cfg.org_dir_title2} {cfg.org_short}"
    merge_block(ws, f"A{R}:B{R+2}", dir_block, S["A58"])
    put(ws, R, 4, cfg.org_dir_name, S["D59"]); R += 2
    put(ws, R, 3, "(подпись)", S["C60"]); put(ws, R, 4, "(инициалы, фамилия)", S["C60"]); R += 2
    put(ws, R, 1, "Главный инженер проекта", S["A62"], c2=2)
    put(ws, R, 4, cfg.gip_name, S["D59"]); R += 1
    put(ws, R, 3, "(подпись)", S["C60"]); put(ws, R, 4, "(инициалы, фамилия)", S["C60"]); R += 1
    put(ws, R, 1, '" _____ " ________________ 20__ г.', S["A64"], c2=2, height=12.75); R += 1
    put(ws, R, 1, "М.П.", S["A65"], height=12.75); R += 2
    put(ws, R, 1, "Заказчик:", S["A67"], height=12.75); R += 1
    cust_block = f"{cfg.cust_dir_title} {cfg.cust_dir_title2}"
    merge_block(ws, f"A{R}:B{R+1}", cust_block, S["A68"])
    put(ws, R, 4, cfg.cust_dir_name, S["D59"]); R += 1
    put(ws, R, 3, "(подпись)", S["C60"]); put(ws, R, 4, "(инициалы, фамилия)", S["C60"]); R += 2
    put(ws, R, 3, "(подпись)", S["C60"]); put(ws, R, 4, "(инициалы, фамилия)", S["C60"]); R += 1
    put(ws, R, 1, '" _____ " ________________ 20__ г.', S["A64"], c2=2, height=12.75); R += 1
    put(ws, R, 1, "М.П.", S["A65"], height=12.75)

    ws.print_area = f"A1:E{R}"
    ws.page_setup.orientation = "portrait"
    ws.page_setup.scale = 68
    return title
