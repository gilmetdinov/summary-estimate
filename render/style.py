"""
Низкоуровневые хелперы рендеринга: копирование стилей, запись ячеек с объединением,
допустимые имена файлов/листов, суффиксы приложений.

Здесь нет бизнес-логики смет — только примитивы работы с openpyxl и строками.
"""
from __future__ import annotations
import os
import re
from copy import copy

from openpyxl.utils import range_boundaries
from openpyxl.worksheet.properties import PageSetupProperties

# суффикс номера сметы -> номер приложения в печатной форме
SUFFIX_APP = {"01.01": "1.1", "01.02": "1.2", "02.01": "2.1", "02.02": "2.2", "03": "3"}

# символы, недопустимые в именах файлов Windows
_ILLEGAL_FNAME = re.compile(r'[<>:"/\\|?*\x00-\x1f]')


def cp(donor, tgt):
    """Копирует оформление ячейки-донора в целевую ячейку."""
    tgt.font = copy(donor.font)
    tgt.border = copy(donor.border)
    tgt.fill = copy(donor.fill)
    tgt.alignment = copy(donor.alignment)
    tgt.protection = copy(donor.protection)
    tgt.number_format = donor.number_format


def put(ws, r, c1, value, donor_cell, c2=None, height=None):
    """Пишет значение в (r,c1), копирует стиль донора во все ячейки c1..c2, объединяет."""
    cc2 = c2 or c1
    for c in range(c1, cc2 + 1):
        cp(donor_cell, ws.cell(row=r, column=c))
    ws.cell(row=r, column=c1).value = value
    if c2 and c2 > c1:
        ws.merge_cells(start_row=r, start_column=c1, end_row=r, end_column=c2)
    if height is not None:
        ws.row_dimensions[r].height = height


def merge_block(ws, cell_range, value, donor_cell):
    """Объединяет прямоугольный диапазон ОДИН раз, скопировав стиль донора во все его
    ячейки. Важно для многострочных шапок: повторный merge поверх уже объединённых
    ячеек Excel считает ошибкой («удалены записи: объединение ячеек»)."""
    min_c, min_r, max_c, max_r = range_boundaries(cell_range)
    for r in range(min_r, max_r + 1):
        for c in range(min_c, max_c + 1):
            cp(donor_cell, ws.cell(row=r, column=c))
    ws.cell(row=min_r, column=min_c).value = value
    ws.merge_cells(cell_range)


def appendix(suffix: str) -> str:
    return SUFFIX_APP.get(suffix, suffix.replace("0", "").replace(".", "."))


def clean_quotes(s: str) -> str:
    """Заменяет кавычки-ёлочки «» на прямые \"\"."""
    if not s:
        return s
    return s.replace("\u00ab", '"').replace("\u00bb", '"')


def safe_filename(name: str) -> str:
    name = _ILLEGAL_FNAME.sub("", name)
    name = re.sub(r"\s+", " ", name).strip().strip(". ")
    return name or "Сводник"


def sheet_title(base: str, used: set) -> str:
    """Уникальное имя листа (не длиннее 31 символа, Excel-ограничение)."""
    base = base[:31]
    t, i = base, 1
    while t in used:
        i += 1
        suffix = f"_{i}"
        t = f"{base[:31 - len(suffix)]}{suffix}"
    used.add(t)
    return t


def unique_name(name: str, used: set) -> str:
    """Уникальное имя файла при батч-конвертации."""
    if name not in used:
        used.add(name)
        return name
    stem, ext = os.path.splitext(name)
    i = 2
    while f"{stem}_{i}{ext}" in used:
        i += 1
    out = f"{stem}_{i}{ext}"
    used.add(out)
    return out


def apply_page_setup(ws, donor_ws, fit_width=False):
    """Наследует печатные настройки с листа-донора (точечное наследование).

    Копирует ориентацию, формат бумаги, поля и сквозные строки/столбцы —
    лист, собранный с нуля, иначе их теряет. fit_width=True включает
    вписывание по ширине страницы (как выгрузка ГС) вместо фикс. масштаба.
    """
    src, dst = donor_ws.page_setup, ws.page_setup
    dst.orientation = src.orientation or "portrait"
    if src.paperSize:
        dst.paperSize = src.paperSize
    m_src, m_dst = donor_ws.page_margins, ws.page_margins
    for attr in ("left", "right", "top", "bottom", "header", "footer"):
        setattr(m_dst, attr, getattr(m_src, attr))
    if donor_ws.print_title_rows:
        ws.print_title_rows = donor_ws.print_title_rows
    if donor_ws.print_title_cols:
        ws.print_title_cols = donor_ws.print_title_cols
    if fit_width:
        dst.fitToWidth = 1
        dst.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    elif src.scale:
        dst.scale = src.scale
