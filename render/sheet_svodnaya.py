"""
Рендер сводной сметы: шапка, строки по сметам (со ссылками на листы), итоги с НДС, подписи.
"""
from __future__ import annotations
from copy import copy

from openpyxl.worksheet.properties import PageSetupProperties

from smeta_parser import Smeta
from config import Config
from render.donor import Donor
from render.style import clean_quotes, merge_block, put


def render_svodnaya(wb, smetas: list[Smeta], refs: list[dict], d: Donor, cfg: Config):
    """Сводная смета по проекту со ссылками на листы отдельных смет.

    После строк по сметам (из выгрузки) опционально добавляет строки изысканий
    (НЗ_ИГДИ) и госэкспертизы из конфига — они вставляются вручную, не из ГС.
    """
    ws = wb.create_sheet(title="св.смета", index=0)
    V = d.sv
    _svod_cols = {"A": 5.29, "B": 20.57, "C": 52.0, "D": 17.0,
                  "E": 13.43, "F": 13.5, "G": 13.5,
                  "H": 9.43, "I": 9.14, "L": 71.86, "M": 130.86, "N": 9.14}
    for col, w in _svod_cols.items():
        ws.column_dimensions[col].width = w
    for col in ("L", "M"):
        ws.column_dimensions[col].hidden = True
    ws.sheet_view.showGridLines = False

    proj = smetas[0].project_code
    obj = clean_quotes(smetas[0].object_base)
    org = clean_quotes(smetas[0].org)
    cust = clean_quotes(smetas[0].customer)

    # служебная колонка H: параметры для CONCATENATE (строка: int режет ведущие нули)
    ws.cell(row=1, column=8).value = proj
    ws.cell(row=7, column=8).value = proj

    put(ws, 1, 7, '=CONCATENATE("Приложение №3.",H1)', V["G1"])
    put(ws, 2, 7, "к договору", V["G1"])
    put(ws, 3, 7, "от ____.____._______г.", V["G1"])
    put(ws, 4, 7, "№_________________", V["G1"])
    for r in (2, 3, 4):
        c = ws.cell(row=r, column=7)
        a = copy(c.alignment)
        a.wrap_text = False
        c.alignment = a
    put(ws, 7, 1, '=CONCATENATE("СВОДНАЯ СМЕТА №",H7)', V["A7"], c2=7)
    put(ws, 8, 1, "на проектные и изыскательские работы", V["A8"], c2=7)
    put(ws, 10, 1, "Наименование строительства и стадии проектирования:", V["A10"], c2=7)
    obj_text = f'Проектные и изыскательские работы по объекту "{obj}"'
    put(ws, 11, 1, obj_text, V["A11"], c2=7, height=31.5)
    # дубликат в скрытой колонке M (как в target)
    ws.cell(row=11, column=13).value = obj_text
    put(ws, 13, 1, f"Наименование проектной организации - генерального проектировщика: {org}", V["A13"], c2=7)
    put(ws, 14, 1, f"Наименование организации заказчика: {cust}", V["A14"], c2=7)

    # шапка таблицы. Двустрочные ячейки объединяем одним блоком (без наложения merge).
    merge_block(ws, "A16:A17", "№ п/п", V["A16"])
    ws.row_dimensions[16].height = 12.75
    merge_block(ws, "B16:C17", "Наименования смет на проектные работы и инженерные изыскания, затрат", V["B16"])
    merge_block(ws, "D16:D17", "Обоснование", V["D16"])
    put(ws, 16, 5, "Сметная стоимость, тыс.руб.", V["E16"], c2=7)
    put(ws, 17, 5, "инженерных изысканий", V["E17"], height=25.5)
    put(ws, 17, 6, "проектных работ", V["F17"])
    put(ws, 17, 7, "всего", V["G17"])
    for c, n in [(1, 1), (2, 2), (4, 3), (5, 4), (6, 5), (7, 6)]:
        put(ws, 18, c, n, V["A18"], height=12.75)
    ws.merge_cells("B18:C18")

    # --- список строк сводной: сметы из выгрузки + ручные плейсхолдеры ---
    entries = []
    for sm, rf in zip(smetas, refs):
        izysk = sm.suffix == "03"
        entries.append({
            "name": sm.kind, "osn": f"Смета №{sm.number}",
            "value": f"='{rf['title']}'!{rf['subtotal_addr']}",
            "column": "izysk" if izysk else "proj",
        })
    for ph, tag in ((cfg.izyskaniya, "izysk"), (cfg.gosexpertiza, "gosexp")):
        if ph:
            entries.append({
                "name": ph.name, "osn": ph.osn,
                "value": ph.value, "column": ph.column,
                "_placeholder": tag,
            })

    # строка-итог по объекту (диапазоны охватывают и сметы, и плейсхолдеры)
    row_proj = 19
    first = 20
    last = first + len(entries) - 1
    # H19 = сумма без госэкспертизы; I19 = G19 - H19 (экспертиза)
    gosexp_row = last if cfg.gosexpertiza else None
    h19_formula = f"=G{row_proj}-G{gosexp_row}" if gosexp_row else f"=G{row_proj}"
    put(ws, row_proj, 1, 1, V["A19"], height=47.25)
    put(ws, row_proj, 2, obj, V["B19"], c2=3)
    put(ws, row_proj, 5, f"=SUM(E{first}:E{last})", V["E19"])
    put(ws, row_proj, 6, f"=SUM(F{first}:F{last})", V["E19"])
    put(ws, row_proj, 7, f"=SUM(G{first}:G{last})", V["E19"])
    put(ws, row_proj, 8, h19_formula, V["H19"])
    put(ws, row_proj, 9, f"=G{row_proj}-H{row_proj}", V["I19"])
    ws.cell(row=row_proj, column=12).value = obj  # L — дубликат B

    R = first
    first_proj = True
    for i, en in enumerate(entries, start=1):
        put(ws, R, 1, f"1.{i}", V["A20"], height=31.5)
        put(ws, R, 2, en["name"], V["B20"], c2=3)
        put(ws, R, 4, en["osn"], V["D20"])
        if en["column"] == "izysk":
            put(ws, R, 5, en["value"], V["E20"]); put(ws, R, 6, None, V["F20"])
        else:
            put(ws, R, 5, None, V["E20"]); put(ws, R, 6, en["value"], V["F20"])
        put(ws, R, 7, f"=SUM(E{R}:F{R})", V["G20"])
        ws.cell(row=R, column=12).value = en["name"]  # L — дубликат B, скрытая
        # O/P: ед. стоимость и стоимость за единицу (всегда, даже без ввода).
        # Строка ед. стоимости проектных работ — первая строка именно проектных
        # работ (не изысканий/экспертизы), а не просто первый entry списка.
        if first_proj and en.get("_placeholder") is None and en["column"] == "proj":
            ws.cell(row=R, column=15).value = cfg.unit_cost_proj
            ws.cell(row=R, column=16).value = f"=F{R}/O{R}*1000"
            first_proj = False
        elif en.get("_placeholder") == "izysk":
            ws.cell(row=R, column=15).value = cfg.unit_cost_izysk
            ws.cell(row=R, column=16).value = f"=E{R}/O{R}*1000"
        elif en.get("_placeholder") == "gosexp":
            ws.cell(row=R, column=15).value = cfg.unit_cost_gosexp
            ws.cell(row=R, column=16).value = f"=F{R}/O{R}*1000"
        R += 1

    put(ws, R, 1, "Всего без НДС", V["A26"], c2=4)
    put(ws, R, 5, None, V["E26"]); put(ws, R, 6, None, V["F26"])
    put(ws, R, 7, f"=G{row_proj}", V["G28"])
    r_bez = R; R += 1
    put(ws, R, 1, f"НДС {cfg.vat_pct}%", V["A27"], c2=4)
    put(ws, R, 5, None, V["E27"]); put(ws, R, 6, None, V["F27"])
    put(ws, R, 7, f"=ROUND(G{r_bez}*{cfg.vat_rate},5)", V["G27"])
    r_nds = R; R += 1
    put(ws, R, 1, "Всего с НДС", V["A28"], c2=4)
    put(ws, R, 5, None, V["E28"]); put(ws, R, 6, None, V["F28"])
    g28_addr = f"G{R}"
    put(ws, R, 7, f"=G{r_bez}+G{r_nds}", V["G28"])
    R += 2

    put(ws, R, 1, "Всего по смете с НДС:", V["A30"], c2=2)
    put(ws, R, 3, d.propis_for_svod(g28_addr), V["C30"])
    R += 3

    # подписи
    put(ws, R, 2, "Проектная организация:", V["B33"]); R += 1
    put(ws, R, 2, cfg.org_dir_title1, V["B34"], c2=3, height=15.75)
    put(ws, R, 6, cfg.org_dir_name, V["F34"], c2=7); R += 1
    put(ws, R, 2, f"{cfg.org_dir_title2} {cfg.org_short}", V["B35"], c2=3)
    put(ws, R, 4, "(подпись)", V["D35"], c2=5); R += 3
    put(ws, R, 2, "Заказчик:", V["B38"]); R += 1
    put(ws, R, 2, cfg.cust_dir_title, V["B39"], c2=3)
    put(ws, R, 6, cfg.cust_dir_name, V["F39"], c2=7); R += 1
    put(ws, R, 2, cfg.cust_dir_title2, V["B39"], c2=3)
    put(ws, R, 4, "(подпись)", V["D35"], c2=5)

    ws.print_area = f"A1:S{R}"
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    return ws
