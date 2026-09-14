"""
Парсер смет на ПИР из «сырой» выгрузки (raw.xlsx) в структурную модель.

Источник — выгрузка из сметной программы (Гранд Смета / Адепт), формат стабильный:
шапка -> таблица (№пп | Наименование | Обоснование МНЗ | Расчёт | Сметная стоимость)
-> позиции с подстроками (коэффициенты + разбивка ПЗ/ППО/ТКР/...) -> итоги -> подписи.

Парсер НЕ привязывается к жёстким номерам строк: он классифицирует строки по содержимому,
поэтому переживает разное число позиций и подстрок.
"""
from __future__ import annotations
import re
from dataclasses import dataclass, field

import openpyxl


def _num(v):
    """'4,08807' | 4.08 | '1\xa0101,2819' -> float; иначе None."""
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace("\xa0", "").replace(" ", "")
    s = s.replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return None


def _txt(v):
    return None if v is None else str(v).strip()


# суффикс номера сметы -> человекочитаемый вид работ (для колонки сводной)
SUFFIX_KIND = {
    # 2-level numbering (example2: "01.01" → suffix "01")
    "01": "Проектные работы (проектная документация)",
    "02": "Проектные работы (рабочая документация)",
    # 3-level numbering (original: "04.01.01" → suffix "01.01")
    "01.01": "Проектные работы (проектная документация)",
    "01.02": "Проектные работы (рабочая документация)",
    "02.01": "Проектные работы (проектная документация)",
    "02.02": "Проектные работы (рабочая документация)",
    "03":    "Инженерно-геодезические изыскания",
}

_TOTAL_START = re.compile(r"^\s*Итог", re.IGNORECASE)
_SECTION = re.compile(r"^\s*Раздел\b", re.IGNORECASE)


@dataclass
class SubLine:
    text: str                 # колонка C (коэффициент-примечание или разбивка)
    value: float | None       # колонка E как число, если есть (разбивка ПЗ/ППО/...)
    value_text: str | None = None  # колонка E как в оригинале (с запятой), для рендера


@dataclass
class Item:
    n: str               # № пп
    name: str            # B
    osn: str             # C — обоснование МНЗ
    calc: str            # D — расчёт стоимости (многострочный)
    total: float | None  # E — сметная стоимость позиции
    sublines: list[SubLine] = field(default_factory=list)


@dataclass
class Section:
    title: str
    items: list[Item] = field(default_factory=list)


@dataclass
class TotalLine:
    label: str
    value: float | None


@dataclass
class Smeta:
    sheet: str
    project_code: str        # '12'
    suffix: str              # '01.01'
    number: str              # '12.01.01'
    kind: str                # вид работ (из суффикса)
    object_full: str         # полное наименование стройки (A6)
    object_base: str         # без хвоста ", Проектные работы (...)"
    work_line: str           # 'на проектные работы' (A5)
    customer: str            # B10
    org: str                 # B13
    price_level: str         # A16
    sections: list[Section] = field(default_factory=list)
    totals: list[TotalLine] = field(default_factory=list)
    subtotal_no_vat: float | None = None  # «Итого Поз. 1-N» (на это ссылается сводная)
    vat: float | None = None
    grand_total: float | None = None


def _split_object(name: str):
    """Отделяет базовое наименование стройки от хвоста с видом работ."""
    if not name:
        return name, name
    for sep in [", Проектные работы", ". Проектные работы",
                ", Инженерно", ". Инженерно", ", Рабочая", ". Рабочая"]:
        i = name.find(sep)
        if i != -1:
            return name[:i].rstrip(" ,."), name
    return name, name


def parse_sheet(ws) -> Smeta:
    g = lambda r, c: ws.cell(row=r, column=c).value

    # --- шапка: ищем по содержимому в колонке A/E, не по фикс. номерам строк ---
    smeta_no = None
    work_line = None
    object_full = None
    customer = None
    org = None
    price_level = None
    header_row = None  # строка с '№ пп'

    for r in range(1, min(ws.max_row, 30) + 1):
        a = _txt(g(r, 1))
        b = _txt(g(r, 2))
        if a:
            m = re.match(r"СМЕТА\s*№\s*([0-9.]+)", a)
            if m:
                smeta_no = m.group(1)
            elif a.lower().startswith("на ") and work_line is None:
                work_line = a
            elif a.startswith("Составлена в уровне цен"):
                price_level = a
            elif a == "№ пп" or a.startswith("№ пп"):
                header_row = r
        # наименование стройки: первая длинная строка в A после номера сметы
        if object_full is None and a and len(a) > 60 and "СМЕТА" not in a:
            object_full = a
        # заказчик/организация лежат в B под метками в A
        if a == "Заказчик" and b is None:
            customer = _txt(g(r + 1, 2))
        if a == "Проектная организация" and b is None:
            org = _txt(g(r + 1, 2))

    if header_row is None:
        raise ValueError(f"[{ws.title}] не найдена шапка таблицы (№ пп)")

    # код проекта и суффикс
    project_code, suffix = "", ""
    if smeta_no:
        parts = smeta_no.split(".", 1)
        project_code = parts[0]
        suffix = parts[1] if len(parts) > 1 else ""
    kind = SUFFIX_KIND.get(suffix, "")
    object_base, _ = _split_object(object_full or "")

    sm = Smeta(
        sheet=ws.title, project_code=project_code, suffix=suffix,
        number=smeta_no or "", kind=kind,
        object_full=object_full or "", object_base=object_base,
        work_line=work_line or "на проектные работы",
        customer=customer or "", org=org or "", price_level=price_level or "",
    )

    # --- тело таблицы: со строки header_row+2 (после строки нумерации колонок) ---
    r = header_row + 1
    # строка нумерации колонок (1 2 3 4 5) — пропускаем
    if str(g(r, 1)).strip() in ("1", "1.0"):
        r += 1

    cur_section: Section | None = None
    cur_item: Item | None = None
    in_totals = False

    while r <= ws.max_row:
        a = _txt(g(r, 1))
        b = _txt(g(r, 2))
        c = _txt(g(r, 3))
        e = g(r, 5)

        # начало блока итогов
        if (b and _TOTAL_START.match(b)) or (a and _TOTAL_START.match(a)):
            in_totals = True

        if in_totals:
            label = b or a
            if label:
                sm.totals.append(TotalLine(label=label, value=_num(e)))
            # подписи начинаются — стоп
            if label and ("Руководитель" in label or "Заказчик" == label):
                break
            r += 1
            continue

        if a and _SECTION.match(a):
            cur_section = Section(title=a)
            sm.sections.append(cur_section)
            cur_item = None
            r += 1
            continue

        # позиция: A — номер, B — наименование, E — число
        if a and re.fullmatch(r"\d+", str(a)) and b:
            if cur_section is None:
                cur_section = Section(title="Раздел 1")
                sm.sections.append(cur_section)
            cur_item = Item(n=a, name=b, osn=c or "",
                            calc=_txt(g(r, 4)) or "", total=_num(e))
            cur_section.items.append(cur_item)
            r += 1
            continue

        # подстрока позиции (только C, иногда + E)
        if cur_item is not None and c:
            ev = _num(e)
            cur_item.sublines.append(
                SubLine(text=c, value=ev, value_text=(_txt(e) if ev is not None else None)))
            r += 1
            continue

        r += 1

    # --- разбор итогов: без НДС / НДС / всего ---
    for t in sm.totals:
        lab = t.label.lower()
        if "ндс" in lab and t.value is not None:
            sm.vat = t.value
        elif "всего по смете" in lab and t.value is not None:
            sm.grand_total = t.value
        elif "итого поз" in lab and t.value is not None and sm.subtotal_no_vat is None:
            sm.subtotal_no_vat = t.value
    if sm.subtotal_no_vat is None:
        # запасной вариант: «Итого по разделу» / «Итого»
        for t in sm.totals:
            if t.value is not None and "ндс" not in t.label.lower() \
               and "всего" not in t.label.lower():
                sm.subtotal_no_vat = t.value
                break
    return sm


def parse_workbook(path: str) -> list[Smeta]:
    wb = openpyxl.load_workbook(path, data_only=True)
    out = []
    for ws in wb.worksheets:
        try:
            out.append(parse_sheet(ws))
        except Exception as ex:  # noqa
            print(f"!! пропущен лист {ws.title!r}: {ex}")
    wb.close()
    return out


if __name__ == "__main__":
    import sys
    path = sys.argv[1] if len(sys.argv) > 1 else "raw.xlsx"
    smetas = parse_workbook(path)
    print(f"Распарсено смет: {len(smetas)}\n")
    for sm in smetas:
        print("=" * 80)
        print(f"№ {sm.number}  [{sm.suffix}] {sm.kind}")
        print(f"  объект(base): {sm.object_base[:70]}...")
        print(f"  заказчик: {sm.customer}")
        print(f"  организация: {sm.org}")
        print(f"  уровень цен: {sm.price_level}")
        for sec in sm.sections:
            print(f"  -- {sec.title}  (позиций: {len(sec.items)})")
            for it in sec.items:
                print(f"     [{it.n}] сумма={it.total}  подстрок={len(it.sublines)}")
                print(f"         наим: {it.name[:60]}...")
                for sl in it.sublines:
                    mark = f" = {sl.value}" if sl.value is not None else ""
                    print(f"           · {sl.text[:60]}{mark}")
        print(f"  ИТОГ без НДС={sm.subtotal_no_vat}  НДС={sm.vat}  ВСЕГО={sm.grand_total}")
