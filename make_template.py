"""
Сборка мастер-шаблона template.xlsx из эталона example.xlsx.

Мастер-шаблон — выделенный файл-донор ОФОРМЛЕНИЯ: только два листа-донора (лист
сметы и лист сводной) под стабильными именами. example.xlsx остаётся эталоном/фикстурой
и от прода отвязан. Рендерер (smeta_render.py) и сборка exe ссылаются на template.xlsx.

Листы копируются «как есть» (тот же файл, лишние листы удалены) — поэтому стили, рамки,
форматы, высоты строк и формулы «сумма прописью» сохраняются с максимальным фиделити.

Запуск:  python3 make_template.py      # перегенерировать из example.xlsx
"""
from __future__ import annotations
import openpyxl

SRC = "sources/example.xlsx"   # эталон-источник (локальный, под .gitignore)
DST = "template.xlsx"          # мастер-шаблон (коммитится, вкладывается в exe)

# исходные листы-доноры в example.xlsx
SRC_SMETA = "смета 04.01.01_ПД_без БИМ - Фор"
SRC_SVOD = "св.смета_04"

# стабильные имена в мастер-шаблоне (ДОЛЖНЫ совпадать с константами в smeta_render.py)
DONOR_SMETA = "ШАБЛОН_смета"
DONOR_SVOD = "ШАБЛОН_сводная"


def build(src: str = SRC, dst: str = DST) -> str:
    wb = openpyxl.load_workbook(src)
    keep = {SRC_SMETA, SRC_SVOD}
    missing = keep - set(wb.sheetnames)
    if missing:
        raise ValueError(f"В {src!r} нет листов-доноров: {missing}")
    for name in list(wb.sheetnames):
        if name not in keep:
            del wb[name]
    wb[SRC_SMETA].title = DONOR_SMETA
    wb[SRC_SVOD].title = DONOR_SVOD
    wb.save(dst)
    print(f"OK -> {dst}  (листы: {wb.sheetnames})")
    return dst


if __name__ == "__main__":
    build()
